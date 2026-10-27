"""文档解析与可定位文本分块。"""

from __future__ import annotations

from io import BytesIO
import hashlib
from pathlib import Path
import re
from typing import Iterable

from docx import Document
from pypdf import PdfReader

from .models import DocumentChunk, ParsedDocument


class DocumentParseError(ValueError):
    """文件无法安全解析时抛出。"""


def decode_text_bytes(data: bytes) -> str:
    if not data:
        return ""
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "utf-16"):
        try:
            text = data.decode(encoding)
            if "\x00" not in text:
                return text.replace("\r\n", "\n").replace("\r", "\n")
        except UnicodeDecodeError:
            continue
    raise DocumentParseError("无法识别文本编码，请将文件另存为 UTF-8 后重试。")


def _source_prefix(source_name: str) -> str:
    return hashlib.sha1(source_name.encode("utf-8")).hexdigest()[:8]


def _split_oversized_block(block: str, max_chars: int) -> list[str]:
    if len(block) <= max_chars:
        return [block]
    sentences = [part.strip() for part in re.split(r"(?<=[。！？!?；;])", block) if part.strip()]
    if not sentences:
        return [block[index : index + max_chars] for index in range(0, len(block), max_chars)]
    pieces: list[str] = []
    buffer = ""
    for sentence in sentences:
        if len(sentence) > max_chars:
            if buffer:
                pieces.append(buffer)
                buffer = ""
            pieces.extend(sentence[index : index + max_chars] for index in range(0, len(sentence), max_chars))
        elif len(buffer) + len(sentence) <= max_chars:
            buffer += sentence
        else:
            if buffer:
                pieces.append(buffer)
            buffer = sentence
    if buffer:
        pieces.append(buffer)
    return pieces


def split_text_blocks(text: str, max_chars: int = 700) -> list[str]:
    normalized = text.replace("\r\n", "\n").replace("\r", "\n").strip()
    if not normalized:
        return []
    raw_blocks = re.split(r"\n\s*\n|\n", normalized)
    blocks: list[str] = []
    for raw in raw_blocks:
        block = re.sub(r"[ \t]+", " ", raw).strip()
        if not block:
            continue
        blocks.extend(_split_oversized_block(block, max_chars=max_chars))
    return blocks


def _make_chunks(
    *,
    source_name: str,
    role: str,
    blocks: Iterable[tuple[str, str, int | None, int | None]],
    is_synthetic: bool,
) -> list[DocumentChunk]:
    prefix = _source_prefix(source_name)
    chunks: list[DocumentChunk] = []
    for index, (text, location, page, paragraph) in enumerate(blocks, start=1):
        clean = text.strip()
        if not clean:
            continue
        chunks.append(
            DocumentChunk(
                chunk_id=f"{prefix}-{index:03d}",
                source_name=source_name,
                role=role,
                location=location,
                text=clean,
                page=page,
                paragraph=paragraph,
                is_synthetic=is_synthetic,
            )
        )
    return chunks


def _parse_txt(name: str, role: str, data: bytes, is_synthetic: bool) -> ParsedDocument:
    text = decode_text_bytes(data)
    blocks = []
    for index, block in enumerate(split_text_blocks(text), start=1):
        blocks.append((block, f"段落 {index}", None, index))
    warnings = [] if blocks else ["文本文件未检测到可提取内容。"]
    return ParsedDocument(
        name=name,
        role=role,
        content_type="text",
        chunks=_make_chunks(source_name=name, role=role, blocks=blocks, is_synthetic=is_synthetic),
        warnings=warnings,
        is_synthetic=is_synthetic,
        metadata={"file_size": len(data), "character_count": len(text)},
    )


def _parse_pdf(name: str, role: str, data: bytes, is_synthetic: bool) -> ParsedDocument:
    try:
        reader = PdfReader(BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as exc:  # pragma: no cover - depends on encrypted PDF implementation
                raise DocumentParseError("PDF 已加密，当前版本无法读取，请提供未加密版本。") from exc
        blocks: list[tuple[str, str, int | None, int | None]] = []
        warnings: list[str] = []
        for page_number, page in enumerate(reader.pages, start=1):
            try:
                page_text = page.extract_text() or ""
            except Exception:
                page_text = ""
            page_blocks = split_text_blocks(page_text)
            if not page_blocks:
                warnings.append(f"第 {page_number} 页未提取到文本，可能是扫描页或图片页。")
                continue
            for paragraph_number, block in enumerate(page_blocks, start=1):
                location = f"第 {page_number} 页 · 段落 {paragraph_number}"
                blocks.append((block, location, page_number, paragraph_number))
    except DocumentParseError:
        raise
    except Exception as exc:
        raise DocumentParseError(f"PDF 解析失败：{exc}") from exc

    if not blocks:
        warnings.append("PDF 中没有可用于证据检索的文本。请提供可复制文本版本，或人工转写关键内容。")
    return ParsedDocument(
        name=name,
        role=role,
        content_type="application/pdf",
        chunks=_make_chunks(source_name=name, role=role, blocks=blocks, is_synthetic=is_synthetic),
        warnings=warnings,
        is_synthetic=is_synthetic,
        metadata={"file_size": len(data), "page_count": len(reader.pages)},
    )


def _parse_docx(name: str, role: str, data: bytes, is_synthetic: bool) -> ParsedDocument:
    try:
        document = Document(BytesIO(data))
    except Exception as exc:
        raise DocumentParseError(f"DOCX 解析失败：{exc}") from exc

    blocks: list[tuple[str, str, int | None, int | None]] = []
    paragraph_index = 0
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        paragraph_index += 1
        for piece_index, piece in enumerate(split_text_blocks(text), start=1):
            suffix = f" · 分段 {piece_index}" if piece_index > 1 else ""
            blocks.append((piece, f"段落 {paragraph_index}{suffix}", None, paragraph_index))

    for table_index, table in enumerate(document.tables, start=1):
        rows: list[str] = []
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        for row_index, row_text in enumerate(rows, start=1):
            blocks.append((row_text, f"表格 {table_index} · 第 {row_index} 行", None, None))

    warnings = [] if blocks else ["DOCX 中未检测到可提取的段落或表格文本。"]
    return ParsedDocument(
        name=name,
        role=role,
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        chunks=_make_chunks(source_name=name, role=role, blocks=blocks, is_synthetic=is_synthetic),
        warnings=warnings,
        is_synthetic=is_synthetic,
        metadata={"file_size": len(data), "paragraph_count": len(document.paragraphs), "table_count": len(document.tables)},
    )


def parse_document_bytes(
    *,
    name: str,
    role: str,
    data: bytes,
    is_synthetic: bool = False,
) -> ParsedDocument:
    suffix = Path(name).suffix.lower().lstrip(".")
    if suffix in {"txt", "md"}:
        doc = _parse_txt(name, role, data, is_synthetic)
    elif suffix == "pdf":
        doc = _parse_pdf(name, role, data, is_synthetic)
    elif suffix == "docx":
        doc = _parse_docx(name, role, data, is_synthetic)
    elif suffix in {"png", "jpg", "jpeg"}:
        raise DocumentParseError("图片不进行离线文字猜测。请上传可提取文字的文档，或在界面中填写海报人工转写。")
    else:
        raise DocumentParseError(f"不支持的文件格式：.{suffix or '未知'}。支持 TXT、MD、PDF、DOCX。")
    return doc


def parse_text_document(
    *,
    name: str,
    role: str,
    text: str,
    is_synthetic: bool = False,
    content_type: str = "text",
) -> ParsedDocument:
    blocks = []
    for index, block in enumerate(split_text_blocks(text), start=1):
        blocks.append((block, f"人工转写 · 段落 {index}", None, index))
    return ParsedDocument(
        name=name,
        role=role,
        content_type=content_type,
        chunks=_make_chunks(source_name=name, role=role, blocks=blocks, is_synthetic=is_synthetic),
        warnings=[] if blocks else ["人工转写内容为空。"],
        is_synthetic=is_synthetic,
        metadata={"character_count": len(text)},
    )
