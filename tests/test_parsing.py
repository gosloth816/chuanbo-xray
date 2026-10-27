from io import BytesIO

from docx import Document

from src.parsing import decode_text_bytes, parse_document_bytes, parse_text_document, split_text_blocks


def test_decode_utf8_and_gb18030():
    assert decode_text_bytes("中文测试".encode("utf-8")) == "中文测试"
    assert decode_text_bytes("中文测试".encode("gb18030")) == "中文测试"


def test_split_text_blocks_respects_limit():
    text = "第一句。" * 200
    blocks = split_text_blocks(text, max_chars=120)
    assert blocks
    assert all(len(block) <= 120 for block in blocks)


def test_parse_text_document_has_locations():
    doc = parse_text_document(name="宣传文案.txt", role="promotion", text="第一段。\n\n第二段。")
    assert len(doc.chunks) == 2
    assert doc.chunks[0].location == "人工转写 · 段落 1"
    assert doc.chunks[1].role == "promotion"


def test_parse_docx():
    document = Document()
    document.add_paragraph("试验结果为防效 82.4%。")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "作物"
    table.cell(0, 1).text = "番茄"
    buffer = BytesIO()
    document.save(buffer)
    parsed = parse_document_bytes(name="资料.docx", role="reference", data=buffer.getvalue())
    assert parsed.chunks
    assert any("82.4%" in chunk.text for chunk in parsed.chunks)
    assert any(chunk.location.startswith("表格 1") for chunk in parsed.chunks)
